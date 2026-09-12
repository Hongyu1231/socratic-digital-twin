"""Synthetic-only tests for the bounded expert interview importer."""

from __future__ import annotations

import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
import uuid
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


SPEC = importlib.util.spec_from_file_location(
    "expert_interview_import",
    Path(__file__).with_name("import-expert-interview.py"),
)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class ExpertInterviewImportTests(unittest.TestCase):
    def make_document(self, paragraphs: list[tuple[int | None, str]], external_link: bool = False) -> bytes:
        document = Document()
        for level, text in paragraphs:
            if level:
                document.add_heading(text, level=level)
            else:
                document.add_paragraph(text)
        if external_link:
            document.part.relate_to("https://example.invalid/not-followed", RT.HYPERLINK, is_external=True)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def make_pack(self, root: Path, case_count: int = 3) -> dict:
        media_id = str(uuid.uuid4())
        media_bytes = b"RIFF\x18\x00\x00\x00WEBPVP8 " + b"synthetic"
        (root / "media").mkdir(parents=True)
        (root / "media" / f"{media_id}.webp").write_bytes(media_bytes)
        cases = []
        for number in range(1, case_count + 1):
            case_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"synthetic-case-{number}"))
            attachments = []
            if number == 1:
                attachments = [{"id": media_id, "kind": "image", "title": "Synthetic OPG", "description": "Synthetic media", "url": "/media"}]
            cases.append({
                "case": {
                    "id": case_id,
                    "title": f"Case {number} - Synthetic case",
                    "description": "Synthetic case description",
                    "attachments": attachments,
                },
                "expertNotes": "Synthetic expert notes",
                "sourceDocument": f"case-{number}.docx",
            })
        manifest = {
            "formatVersion": 1,
            "packageId": "a" * 64,
            "cases": cases,
            "articles": [{
                "id": "existing-source",
                "title": "Existing source",
                "filename": "existing.pdf",
                "sha256": "b" * 64,
                "pages": [{"page": 1, "text": "Existing source text"}],
            }],
            "media": [{
                "id": media_id,
                "caseId": cases[0]["case"]["id"],
                "file": f"media/{media_id}.webp",
                "mimeType": "image/webp",
                "sha256": MODULE.sha256(media_bytes),
                "width": 10,
                "height": 10,
            }],
        }
        (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        return manifest

    def test_sections_experts_and_case_scoped_chunks(self):
        paragraphs = [
            (2, "Experts’ general approach to managing impacted canine cases"),
            (3, "Expert 1"),
            (None, "Interviewer: What do you evaluate first?"),
            (None, "Orthodontist: I evaluate position, space, age, and patient preference."),
            (2, "Comments on Case 1"),
            (3, "Expert 2"),
            (None, "Interviewer: What is your recommendation?"),
            (None, "Orthodontist: Preserve the tooth when the case-specific evidence supports traction."),
            (2, "Comments on Case 2"),
            (3, "Expert 3"),
            (None, "Oral surgeon: What is the prognosis?"),
            (None, "Clinician: It is guarded, and the answer depends on the three-dimensional position."),
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "pack"
            root.mkdir()
            manifest = self.make_pack(root)
            doc = Path(directory) / "panel.docx"
            doc.write_bytes(self.make_document(paragraphs))
            output = Path(directory) / "revision"
            report = MODULE.build(doc, root, output)
            result = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
            self.assertFalse(report["noop"])
            self.assertEqual(result["packageId"], MODULE.revision_package_id("a" * 64, MODULE.sha256(doc.read_bytes())))
            self.assertEqual(len(result["cases"]), len(manifest["cases"]))
            article = next(item for item in result["articles"] if item.get("sourceType") == "expert_interview")
            self.assertRegex(article["id"], MODULE.UUID_RE)
            self.assertEqual(article["sha256"], MODULE.sha256(doc.read_bytes()))
            self.assertEqual(article["sourceHash"], article["sha256"])
            self.assertTrue(article["pages"])
            self.assertTrue(all(len(page["text"]) <= MODULE.MAX_CHUNK_CHARS for page in article["pages"]))
            self.assertTrue(all("DOCX paragraphs" in page["locator"] and "section" in page["locator"] for page in article["pages"]))
            general = [page for page in article["pages"] if page["section"].startswith("Experts’")]
            case_one = [page for page in article["pages"] if page["section"] == "Comments on Case 1"]
            self.assertTrue(general and all("caseIds" not in page for page in general))
            self.assertTrue(case_one and all(page["caseIds"] == [manifest["cases"][0]["case"]["id"]] for page in case_one))
            self.assertTrue(any("Expert 1" in page["text"] for page in general))
            self.assertTrue(any("Interviewer:" in page["text"] and "Orthodontist:" in page["text"] for page in article["pages"]))

    def test_large_paragraph_has_bounded_lossless_chunks_and_provenance(self):
        answer = " ".join(f"sentence-{index} explains the case-specific reasoning." for index in range(500))
        paragraphs = [
            (2, "General approach"),
            (3, "Expert 1"),
            (None, "Interviewer: What is the approach?"),
            (None, answer),
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "pack"
            root.mkdir()
            self.make_pack(root, case_count=1)
            doc = Path(directory) / "long.docx"
            doc.write_bytes(self.make_document(paragraphs))
            output = Path(directory) / "revision"
            MODULE.build(doc, root, output)
            article = next(item for item in json.loads((output / "manifest.json").read_text(encoding="utf-8"))["articles"] if item.get("sourceType") == "expert_interview")
            self.assertGreater(len(article["pages"]), 2)
            joined = " ".join(page["text"] for page in article["pages"])
            for token in ("sentence-0", "sentence-249", "sentence-499"):
                self.assertIn(token, joined)
            self.assertTrue(all(len(page["text"]) <= MODULE.MAX_CHUNK_CHARS for page in article["pages"]))
            self.assertTrue(all("DOCX paragraphs" in page["locator"] for page in article["pages"]))

    def test_identical_source_is_deduplicated_and_repeat_is_noop(self):
        paragraphs = [(2, "General approach"), (3, "Expert 1"), (None, "Interviewer: What matters?"), (None, "Orthodontist: Space matters.")]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "pack"
            root.mkdir()
            self.make_pack(root, case_count=1)
            doc = Path(directory) / "panel.docx"
            doc.write_bytes(self.make_document(paragraphs))
            output = Path(directory) / "revision"
            first = MODULE.build(doc, root, output)
            before = (output / "manifest.json").read_bytes()
            second = MODULE.build(doc, root, output)
            self.assertFalse(first["noop"])
            self.assertTrue(second["noop"])
            self.assertEqual(before, (output / "manifest.json").read_bytes())
            result = json.loads(before)
            source_articles = [item for item in result["articles"] if item.get("sourceHash") == first["sourceHash"]]
            self.assertEqual(len(source_articles), 1)

    def test_unsafe_media_path_and_existing_different_output_are_rejected(self):
        paragraphs = [(2, "General approach"), (3, "Expert 1"), (None, "Interviewer: What matters?"), (None, "Orthodontist: Space matters.")]
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            root = base / "pack"
            root.mkdir()
            manifest = self.make_pack(root, case_count=1)
            manifest["media"][0]["file"] = "../outside.webp"
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            doc = base / "panel.docx"
            doc.write_bytes(self.make_document(paragraphs))
            with self.assertRaises(ValueError):
                MODULE.build(doc, root, base / "revision")

            good_root = base / "good-pack"
            good_root.mkdir()
            self.make_pack(good_root, case_count=1)
            output = base / "revision"
            output.mkdir()
            (output / "manifest.json").write_text(json.dumps({"formatVersion": 1, "packageId": "c" * 64, "cases": [], "articles": [], "media": []}), encoding="utf-8")
            with self.assertRaises(ValueError):
                MODULE.build(doc, good_root, output)

    def test_external_docx_relationship_is_ignored_without_network_following(self):
        paragraphs = [(2, "General approach"), (3, "Expert 1"), (None, "Interviewer: What matters?"), (None, "Orthodontist: Only the supplied record matters.")]
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            root = base / "pack"
            root.mkdir()
            self.make_pack(root, case_count=1)
            doc = base / "panel.docx"
            doc.write_bytes(self.make_document(paragraphs, external_link=True))
            output = base / "revision"
            report = MODULE.build(doc, root, output)
            self.assertGreaterEqual(report["externalLinksIgnored"], 1)
            article = next(item for item in json.loads((output / "manifest.json").read_text(encoding="utf-8"))["articles"] if item.get("sourceType") == "expert_interview")
            self.assertNotIn("example.invalid", json.dumps(article))


if __name__ == "__main__":
    unittest.main()
